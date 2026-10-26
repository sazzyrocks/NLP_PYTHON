"""In-memory document text extraction for PDF and DOCX files."""

import io
from pathlib import Path
from typing import Union
import docx
import fitz  # PyMuPDF


class ExtractionError(Exception):
    """Base exception for document extraction failures."""
    pass


class UnsupportedFormatError(ExtractionError):
    """Raised when the uploaded file type is not supported."""
    pass


class EmptyFileError(ExtractionError):
    """Raised when the uploaded file has 0 bytes."""
    pass


class NoTextExtractedError(ExtractionError):
    """Raised when no text layer can be extracted (e.g. scanned image PDF)."""
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from PDF bytes in memory.

    Args:
        file_bytes: Raw bytes of the PDF file.

    Returns:
        Extracted plain text string.

    Raises:
        ExtractionError: If the PDF cannot be opened or is corrupted.
        NoTextExtractedError: If the PDF contains no extractable text.
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise ExtractionError(f"Failed to read PDF document: {exc}") from exc

    try:
        extracted_pages = []
        for page_index in range(len(doc)):
            page = doc[page_index]
            text = page.get_text("text")
            if text:
                extracted_pages.append(text)
    finally:
        doc.close()

    full_text = "\n".join(extracted_pages).strip()
    if not full_text:
        raise NoTextExtractedError(
            "The PDF contains no extractable text layer. "
            "It may be scanned or image-only."
        )

    return full_text


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from DOCX bytes in memory.

    Args:
        file_bytes: Raw bytes of the DOCX file.

    Returns:
        Extracted plain text string.

    Raises:
        ExtractionError: If the DOCX cannot be opened or parsed.
        NoTextExtractedError: If the DOCX contains no text.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise ExtractionError(f"Failed to read DOCX document: {exc}") from exc

    parts = []
    # Extract paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            parts.append(txt)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                parts.append(" | ".join(row_texts))

    full_text = "\n".join(parts).strip()
    if not full_text:
        raise NoTextExtractedError("The DOCX document contains no readable text.")

    return full_text


def extract_text(filename: Union[str, Path], file_bytes: bytes) -> str:
    """Extract text from an uploaded file based on its extension.

    Supports .pdf and .docx. Operates strictly in-memory without disk writes.

    Args:
        filename: Name of the uploaded file.
        file_bytes: Binary contents of the uploaded file.

    Returns:
        Extracted text.

    Raises:
        EmptyFileError: When file_bytes is empty.
        UnsupportedFormatError: When extension is not .pdf or .docx.
        NoTextExtractedError: When no text layer is found.
        ExtractionError: For parsing and file corruption errors.
    """
    if not file_bytes:
        raise EmptyFileError(f"The file '{filename}' is empty (0 bytes).")

    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_bytes)
    elif ext == ".docx":
        return extract_text_from_docx(file_bytes)
    else:
        raise UnsupportedFormatError(
            f"Unsupported file format '{ext}' for file '{filename}'. "
            f"Only .pdf and .docx files are supported."
        )
